"""Ingestion test fixtures: in-memory sample files per type, an on-disk sample folder,
and (for the pipeline tests) testcontainers Postgres + MinIO with the core stores.
"""

from __future__ import annotations

import io
import os

os.environ.setdefault("TESTCONTAINERS_RYUK_DISABLED", "true")

from collections.abc import AsyncIterator, Callable, Iterator
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import AsyncEngine, AsyncSession, async_sessionmaker
from testcontainers.minio import MinioContainer
from testcontainers.postgres import PostgresContainer

from metis_core.audit import PostgresAuditSink
from metis_core.db.engine import make_engine, make_sessionmaker
from metis_core.dev.testing import (
    async_url,
    docker_available,
    make_minio,
    make_postgres,
    minio_settings,
    run_upgrade,
    truncate_all,
)
from metis_core.objectstore import S3ObjectStore
from metis_core.stores import (
    PostgresClaimStore,
    PostgresDocumentStore,
    PostgresMinioArtifactStore,
)
from metis_ingestion import IngestionPipeline, LocalFolderConnector
from metis_protocol import PolicyState, WorkspaceId

WORKSPACE = WorkspaceId(f"ws_{'a' * 32}")
CONNECTORS_ROOT = Path(__file__).parent.parent / "fixtures" / "connectors"


@pytest.fixture
def workspace() -> WorkspaceId:
    return WORKSPACE


@pytest.fixture
def connectors_root() -> Path:
    """Root of the recorded per-connector fixtures (the credential-free replay corpus)."""
    return CONNECTORS_ROOT


def _make_pdf(lines: list[str]) -> bytes:
    ops = b"BT /F1 12 Tf 72 720 Td 14 TL\n"
    for line in lines:
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        ops += b"(" + escaped.encode("latin-1") + b") Tj T*\n"
    ops += b"ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(ops)).encode() + b" >>\nstream\n" + ops + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += str(index).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    xref = len(pdf)
    pdf += b"xref\n0 " + str(len(objects) + 1).encode() + b"\n0000000000 65535 f \n"
    for offset in offsets:
        pdf += (f"{offset:010d} 00000 n \n").encode()
    pdf += (
        b"trailer\n<< /Size " + str(len(objects) + 1).encode() + b" /Root 1 0 R >>\n"
        b"startxref\n" + str(xref).encode() + b"\n%%EOF"
    )
    return bytes(pdf)


def _make_docx(paragraphs: list[str]) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _make_xlsx(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "People"
    for row in rows:
        worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def samples() -> dict[str, tuple[str, bytes]]:
    """One small sample per supported type: label -> (filename, bytes)."""
    return {
        "txt": ("notes.txt", b"Ada Lovelace is the CTO of Acme Inc.\n\nFounded in 2019."),
        "md": (
            "notes.md",
            b"# Acme\n\nAda Lovelace leads engineering.\n\nGrace Hopper joined in 2020.",
        ),
        "csv": ("people.csv", b"name,role\nAda,CTO\nGrace,Advisor\n"),
        "html": (
            "page.html",
            b"<html><body><h1>Acme</h1><p>Ada Lovelace is the CTO.</p>"
            b"<p>Founded 2019.</p></body></html>",
        ),
        "eml": (
            "mail.eml",
            b"From: ada@acme.com\nTo: team@acme.com\nSubject: Roadmap\n\nWe ship in 2026.",
        ),
        "pdf": (
            "report.pdf",
            _make_pdf(["Ada Lovelace is the CTO of Acme Inc.", "Founded in 2019."]),
        ),
        "docx": (
            "memo.docx",
            _make_docx(["Ada Lovelace is the CTO of Acme Inc.", "Founded 2019."]),
        ),
        "xlsx": (
            "sheet.xlsx",
            _make_xlsx([["name", "role"], ["Ada", "CTO"], ["Grace", "Advisor"]]),
        ),
    }


@pytest.fixture
def sample_dir(tmp_path: Path, samples: dict[str, tuple[str, bytes]]) -> Path:
    for filename, data in samples.values():
        (tmp_path / filename).write_bytes(data)
    return tmp_path


# --- testcontainers fixtures (pipeline integration tests) ---


@pytest.fixture(scope="session")
def postgres() -> Iterator[PostgresContainer]:
    if not docker_available():
        pytest.skip("Docker is not available")
    container = make_postgres()
    container.start()
    try:
        run_upgrade(async_url(container))
        yield container
    finally:
        container.stop()


@pytest.fixture(scope="session")
def minio() -> Iterator[MinioContainer]:
    if not docker_available():
        pytest.skip("Docker is not available")
    container = make_minio()
    container.start()
    try:
        yield container
    finally:
        container.stop()


@pytest.fixture(scope="session")
def object_store(minio: MinioContainer) -> S3ObjectStore:
    import asyncio

    endpoint, access_key, secret_key = minio_settings(minio)
    store = S3ObjectStore(
        bucket="metis-ingest-test",
        endpoint_url=endpoint,
        region="us-east-1",
        access_key=access_key,
        secret_key=secret_key,
    )
    asyncio.run(store.ensure_bucket())
    return store


@pytest.fixture
async def engine(postgres: PostgresContainer) -> AsyncIterator[AsyncEngine]:
    engine = make_engine(async_url(postgres))
    await truncate_all(engine)
    yield engine
    await engine.dispose()


@pytest.fixture
def sessionmaker(engine: AsyncEngine) -> async_sessionmaker[AsyncSession]:
    return make_sessionmaker(engine)


@pytest.fixture
def make_pipeline(
    sessionmaker: async_sessionmaker[AsyncSession], object_store: S3ObjectStore
) -> Callable[[Path], IngestionPipeline]:
    def factory(root: Path) -> IngestionPipeline:
        return IngestionPipeline(
            connector=LocalFolderConnector(root, workspace_id=WORKSPACE, policy=PolicyState()),
            artifact_store=PostgresMinioArtifactStore(sessionmaker, object_store),
            document_store=PostgresDocumentStore(sessionmaker),
            claim_store=PostgresClaimStore(sessionmaker),
            audit_sink=PostgresAuditSink(sessionmaker),
        )

    return factory
