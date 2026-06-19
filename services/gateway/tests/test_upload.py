"""File upload through the gateway (Stage 1 acceptance: every supported format ingests with a
visible parse status). Each upload runs the real parser + extractor; a bad file in a batch is
isolated as a failed status; and the membership gate still guards who may upload into a workspace.
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook


def _provision(client: TestClient, op: dict[str, str], org_id: str, email: str) -> str:
    resp = client.post(
        "/users",
        json={"organization_id": org_id, "email": email, "display_name": email.split("@")[0]},
        headers=op,
    )
    assert resp.status_code == 201, resp.text
    user_id: str = resp.json()["id"]
    return user_id


def _bearer(user_id: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {user_id}"}


def _ada_workspace(client: TestClient, op: dict[str, str]) -> tuple[str, dict[str, str], str]:
    """Provision Ada in a fresh org and return (org_id, Ada's auth, Ada's personal workspace id)."""
    org_id = client.post("/organizations", json={"name": "Acme"}, headers=op).json()["id"]
    ada = _provision(client, op, org_id, "ada@acme.example")
    ws = client.get("/workspaces", headers=_bearer(ada)).json()[0]["id"]
    return org_id, _bearer(ada), ws


def _make_pdf(lines: list[str]) -> bytes:
    ops = b"BT /F1 12 Tf 72 720 Td 14 TL\n"
    for line in lines:
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        ops += b"(" + escaped.encode("latin-1") + b") Tj T*\n"
    ops += b"ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(ops)).encode() + b" >>\nstream\n" + ops + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += str(index).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    xref = len(pdf)
    pdf += b"xref\n0 " + str(len(objects) + 1).encode() + b"\n0000000000 65535 f \n"
    for offset in offsets:
        pdf += (f"{offset:010d} 00000 n \n").encode()
    pdf += (
        b"trailer\n<< /Size " + str(len(objects) + 1).encode() + b" /Root 1 0 R >>\n"
        b"startxref\n" + str(xref).encode() + b"\n%%EOF"
    )
    return bytes(pdf)


def _make_docx(paragraphs: list[str]) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _make_xlsx(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    for row in rows:
        worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# One small sample per supported format — the full PDF/DOCX/XLSX/CSV/TXT/MD/HTML/EML set.
_SAMPLES: dict[str, bytes] = {
    "notes.txt": b"Ada Lovelace is the CTO of Acme Inc.\n\nFounded in 2019.",
    "notes.md": b"# Acme\n\nAda Lovelace leads engineering at Acme.",
    "people.csv": b"name,role\nAda,CTO\nGrace,Advisor\n",
    "page.html": b"<html><body><h1>Acme</h1><p>Ada Lovelace is the CTO.</p></body></html>",
    "mail.eml": b"From: ada@acme.com\nTo: team@acme.com\nSubject: Roadmap\n\nWe ship in 2026.",
    "report.pdf": _make_pdf(["Ada Lovelace is the CTO of Acme Inc.", "Founded in 2019."]),
    "memo.docx": _make_docx(["Ada Lovelace is the CTO of Acme Inc.", "Founded 2019."]),
    "sheet.xlsx": _make_xlsx([["name", "role"], ["Ada", "CTO"], ["Grace", "Advisor"]]),
}


@pytest.mark.parametrize("filename", list(_SAMPLES))
def test_each_supported_format_uploads_with_a_parsed_status(
    client: TestClient, op: dict[str, str], filename: str
) -> None:
    _, ada, ws = _ada_workspace(client, op)
    resp = client.post(
        f"/workspaces/{ws}/upload",
        files=[("files", (filename, _SAMPLES[filename]))],
        headers=ada,
    )
    assert resp.status_code == 201, resp.text
    [status] = resp.json()["files"]
    assert status["status"] == "parsed", status
    assert status["filename"] == filename
    assert status["doc_id"]
    assert status["segments"] >= 1  # the parser produced text; claim yield varies by format


def test_pdf_upload_reports_parse_quality(client: TestClient, op: dict[str, str]) -> None:
    _, ada, ws = _ada_workspace(client, op)
    dense_pdf = _make_pdf(["Ada Lovelace is the CTO of Acme Inc. Founded in 2019."] * 30)
    resp = client.post(
        f"/workspaces/{ws}/upload", files=[("files", ("report.pdf", dense_pdf))], headers=ada
    )
    [status] = resp.json()["files"]
    assert status["page_count"] == 1  # pypdf len(reader.pages), now surfaced
    assert status["parse_path"] == "deterministic"
    assert status["coverage"] is not None
    assert status["coverage"] > 0.5
    assert status["warnings"] == []  # a text-rich PDF has no quality warnings


def test_batch_isolates_a_bad_file_and_still_reports_the_good_one(
    client: TestClient, op: dict[str, str]
) -> None:
    _, ada, ws = _ada_workspace(client, op)
    resp = client.post(
        f"/workspaces/{ws}/upload",
        files=[
            ("files", ("notes.txt", b"Ada Lovelace is the CTO of Acme Inc.")),
            ("files", ("mystery.bin", b"\x00\x01\x02 not a recognized type")),
        ],
        headers=ada,
    )
    assert resp.status_code == 201, resp.text
    by_name = {f["filename"]: f for f in resp.json()["files"]}
    assert by_name["notes.txt"]["status"] == "parsed"
    assert by_name["mystery.bin"]["status"] == "unsupported"  # one bad file did not fail the batch


def test_upload_is_membership_gated(client: TestClient, op: dict[str, str]) -> None:
    org_id, _, ws = _ada_workspace(client, op)
    grace = _provision(client, op, org_id, "grace@acme.example")  # same org, no membership

    resp = client.post(
        f"/workspaces/{ws}/upload",
        files=[("files", ("x.txt", b"nope"))],
        headers=_bearer(grace),
    )
    assert resp.status_code == 403  # the isolation gate guards uploads too


def test_uploaded_evidence_is_queryable(client: TestClient, op: dict[str, str]) -> None:
    _, ada, ws = _ada_workspace(client, op)
    client.post(
        f"/workspaces/{ws}/upload",
        files=[("files", ("acme.txt", b"Ada Lovelace is the CTO of Acme Inc."))],
        headers=ada,
    )
    answered = client.post(
        f"/workspaces/{ws}/query", json={"text": "Who is the CTO of Acme?"}, headers=ada
    )
    assert answered.status_code == 200, answered.text
    assert answered.json()["sufficient"] is True
